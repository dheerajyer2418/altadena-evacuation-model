"""Compile the model's results into a single multi-sheet Excel workbook WITH CHARTS.

    python -m scripts.export_excel        # -> outputs/altadena_results.xlsx

Sheets:
  * READ ME             — plain-language summary + honest caveats
  * Scenario results    — every run's metrics  (+ grouped bar chart)
  * Analysis            — derived shares: reached / carless / not-reached / exposed
                          (+ bar charts), plus an AV-impact comparison
  * Completion over time— % reached vs. time for each scenario (+ line chart)
  * Population fit       — synthesis fit error per neighborhood (+ bar chart)
  * Policy comparison…  — each comparison CSV that exists
  * Data sources        — provenance of every input

Charts are native Excel charts (interactive). Reads whatever result files are
present; missing ones are skipped (never faked).
"""
from __future__ import annotations

import json

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

from altadena_evac.config import project_path

UI = project_path("outputs", "ui")


# --------------------------------------------------------------- static sheets
def _readme_df() -> pd.DataFrame:
    lines = [
        "ALTADENA WILDFIRE EVACUATION MODEL — RESULTS",
        "",
        "PLEASE READ FIRST:",
        "• HYPOTHETICAL 'what-if' scenarios, NOT a reconstruction of the Eaton Fire",
        "  and NOT a prediction. The fire is invented; traffic is NOT calibrated.",
        "• Numbers are meaningful only as a RELATIVE COMPARISON between strategies,",
        "  which all run on the exact same population and fire.",
        "• Software is tested (29 automated checks); real-world accuracy is NOT",
        "  externally validated.",
        "",
        "KEY TERMS:",
        "• Reached            = share of people present who reached safety.",
        "• Carless reached    = same, but only households without a car (the focus).",
        "• Not reached        = share still in transit / waiting / stranded at the end.",
        "• Exposed share      = people the fire front reached / people present (a",
        "                       DANGER signal, NOT a casualty count).",
        "",
        "POPULATION: a Census-based synthetic stand-in, ~15,000 households / ~41,000",
        "people, within ~1% of official ACS totals. Carless households are slightly",
        "overcounted (~4.5% vs ACS 3.7%) — treat that exact number as approximate.",
        "",
        "SEE THE CHARTS on the 'Scenario results', 'Analysis', 'Completion over time',",
        "and 'Population fit' sheets.",
    ]
    return pd.DataFrame({"About these results": lines})


def _data_sources_df() -> pd.DataFrame:
    rows = [
        ("U.S. Census ACS 2019-2023", "Neighborhood counts: age/sex, race, Hispanic, "
         "vehicles, household size, income — 34 Altadena block groups",
         "Who lives where; who has a car"),
        ("ACS PUMS 2023", "Anonymous example households, LA-area PUMAs 03724/03725",
         "Building realistic households"),
        ("Census TIGER/Line 2023", "Place + neighborhood boundaries",
         "Selecting Altadena; placing homes correctly"),
        ("OpenStreetMap", "Drivable roads, Altadena + 3 km buffer (3,766 nodes / "
         "9,972 edges)", "The road network"),
        ("INRIX", "NOT SUPPLIED", "Nothing used; no traffic calibration"),
    ]
    return pd.DataFrame(rows, columns=["Source", "What / geography", "Used for"])


# --------------------------------------------------------------- data from runs
def _index() -> dict:
    p = UI / "index.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def _scenario_results_df() -> pd.DataFrame | None:
    index = _index()
    if not index:
        return None
    rows = []
    for entry in index.values():
        meta_summary = entry.get("summary", {})
        row = {"scenario": entry.get("label", entry.get("id")),
               "population": (entry.get("population", "") or "").split(";")[0]}
        row.update(meta_summary)
        rows.append(row)
    return pd.DataFrame(rows) if rows else None


def _analysis_df() -> pd.DataFrame | None:
    index = _index()
    if not index:
        return None
    rows = []
    for e in index.values():
        s = e.get("summary", {})
        pres = s.get("n_present") or 0
        rows.append({
            "Scenario": e.get("label", e.get("id")),
            "Reached": round(s.get("reached_frac") or 0, 4),
            "Carless reached": round(s.get("carless_reached_frac") or 0, 4),
            "Not reached": round(1 - (s.get("reached_frac") or 0), 4),
            "Exposed share": round((s.get("n_exposed") or 0) / pres, 4) if pres else 0,
            "Trapped vehicles": s.get("n_trapped_vehicles") or 0,
            "Fleet": s.get("fleet_size") or 0,
            "Policy": s.get("policy") or "",
        })
    return pd.DataFrame(rows)


def _completion_df():
    """% reached vs time (minute grid) for each scenario, from the replay frames."""
    index = _index()
    if not index:
        return None, []
    series, labels, maxT = {}, [], 0
    for e in index.values():
        f = UI / e["file"]
        try:
            rep = json.loads(f.read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001
            continue
        pts = []
        for fr in rep.get("frames", []):
            st = fr.get("stats", {})
            pres = st.get("present") or 0
            pts.append((fr.get("t", 0), (st.get("reached", 0) / pres) if pres else 0.0))
        if not pts:
            continue
        label = e.get("label", e.get("id"))
        series[label] = pts
        labels.append(label)
        maxT = max(maxT, pts[-1][0])
    if not labels:
        return None, []
    step = 60
    grid = list(range(0, int(maxT) + step, step))
    data = {"minute": [round(g / 60, 1) for g in grid]}
    for label in labels:
        pts, vals, i, last = series[label], [], 0, 0.0
        for g in grid:
            while i < len(pts) and pts[i][0] <= g:
                last = pts[i][1]
                i += 1
            vals.append(round(last, 4))
        data[label] = vals
    return pd.DataFrame(data), labels


def _population_fit_df() -> pd.DataFrame | None:
    p = project_path("data", "processed", "population_fit_report.json")
    if not p.exists():
        return None
    rep = json.loads(p.read_text(encoding="utf-8"))
    rows = [{"block_group": g, "fit_error_srmse": v.get("srmse"),
             "unsupported_combos": v.get("n_unsupported")}
            for g, v in rep.get("per_block_group_fit", {}).items()]
    df = pd.DataFrame(rows)
    return df.sort_values("fit_error_srmse", ascending=False).reset_index(drop=True) \
        if len(df) else None


# --------------------------------------------------------------- chart helpers
def _bar(ws, df, cats_col, value_cols, title, anchor, y_title=""):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.y_axis.title = y_title
    ch.height, ch.width = 9, 22
    n = len(df)
    for name in value_cols:
        c = df.columns.get_loc(name) + 1
        ch.add_data(Reference(ws, min_col=c, min_row=1, max_row=n + 1),
                    titles_from_data=True)
    cc = df.columns.get_loc(cats_col) + 1
    ch.set_categories(Reference(ws, min_col=cc, min_row=2, max_row=n + 1))
    ws.add_chart(ch, anchor)


def _line(ws, df, x_col, y_cols, title, anchor):
    ch = LineChart()
    ch.title = title
    ch.y_axis.title = "fraction reached safety"
    ch.x_axis.title = "minutes since fire start"
    ch.height, ch.width = 11, 24
    n = len(df)
    for name in y_cols:
        c = df.columns.get_loc(name) + 1
        ch.add_data(Reference(ws, min_col=c, min_row=1, max_row=n + 1),
                    titles_from_data=True)
    xc = df.columns.get_loc(x_col) + 1
    ch.set_categories(Reference(ws, min_col=xc, min_row=2, max_row=n + 1))
    ws.add_chart(ch, anchor)


def _autowidth(xw):
    for ws in xw.book.worksheets:
        for col in ws.columns:
            try:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            except Exception:  # noqa: BLE001
                width = 12
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 60)


def _safe(name: str) -> str:
    for c in '[]:*?/\\':
        name = name.replace(c, " ")
    return name[:31]


def main():
    out = project_path("outputs", "altadena_results.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    comparison_csvs = {
        "Policy cmp (real roads)": project_path("outputs", "compare_altadena",
                                                "policy_comparison_altadena.csv"),
        "Policy cmp (toy)": project_path("outputs", "compare_toy", "policy_comparison.csv"),
    }
    wrote = []
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        _readme_df().to_excel(xw, sheet_name="READ ME", index=False)
        wrote.append("READ ME")

        sr = _scenario_results_df()
        if sr is not None:
            sr.to_excel(xw, sheet_name="Scenario results", index=False)
            ws = xw.sheets["Scenario results"]
            anchor = f"A{len(sr) + 3}"
            cols = [c for c in ["reached_frac", "carless_reached_frac"] if c in sr.columns]
            if "scenario" in sr.columns and cols:
                _bar(ws, sr, "scenario", cols,
                     "Reached vs. carless-reached by scenario", anchor, "fraction")
            wrote.append("Scenario results")

        an = _analysis_df()
        if an is not None:
            an.to_excel(xw, sheet_name="Analysis", index=False)
            ws = xw.sheets["Analysis"]
            base = len(an) + 3
            _bar(ws, an, "Scenario", ["Reached", "Carless reached"],
                 "Who gets out: overall vs. carless", f"A{base}", "fraction")
            _bar(ws, an, "Scenario", ["Not reached", "Exposed share"],
                 "Left behind & exposed (share of people)", f"A{base + 18}", "fraction")
            # AV-impact note for the real-population pair, if present
            labels = list(an["Scenario"])
            base_real = next((i for i, s in enumerate(labels)
                              if "REAL" in s and "No-AV" in s), None)
            av_real = next((i for i, s in enumerate(labels)
                            if "REAL" in s and "No-AV" not in s), None)
            if base_real is not None and av_real is not None:
                r0, c0 = an.iloc[base_real]["Reached"], an.iloc[base_real]["Carless reached"]
                r1, c1 = an.iloc[av_real]["Reached"], an.iloc[av_real]["Carless reached"]
                impact = pd.DataFrame({
                    "AV impact (real population)": [
                        "Carless reached WITHOUT AVs", "Carless reached WITH AVs",
                        "Carless gain (pct points)", "Overall reached WITHOUT AVs",
                        "Overall reached WITH AVs", "Overall gain (pct points)"],
                    "value": [f"{c0:.0%}", f"{c1:.0%}", f"{(c1 - c0) * 100:+.1f}",
                              f"{r0:.0%}", f"{r1:.0%}", f"{(r1 - r0) * 100:+.1f}"]})
                impact.to_excel(xw, sheet_name="Analysis", index=False,
                                startrow=base + 36)
            wrote.append("Analysis")

        comp, labels = _completion_df()
        if comp is not None:
            comp.to_excel(xw, sheet_name="Completion over time", index=False)
            ws = xw.sheets["Completion over time"]
            _line(ws, comp, "minute", labels, "Evacuation completion over time",
                  f"{get_column_letter(len(comp.columns) + 2)}2")
            wrote.append("Completion over time")

        pf = _population_fit_df()
        if pf is not None:
            pf.to_excel(xw, sheet_name="Population fit", index=False)
            ws = xw.sheets["Population fit"]
            _bar(ws, pf, "block_group", ["fit_error_srmse"],
                 "Synthesis fit error by neighborhood (lower = better)",
                 f"E2", "SRMSE")
            wrote.append("Population fit")

        for name, path in comparison_csvs.items():
            if path.exists():
                pd.read_csv(path).to_excel(xw, sheet_name=_safe(name), index=False)
                wrote.append(name)

        _data_sources_df().to_excel(xw, sheet_name="Data sources", index=False)
        wrote.append("Data sources")
        _autowidth(xw)

    print(f"wrote {out}")
    print("sheets:", ", ".join(wrote))


if __name__ == "__main__":
    main()
